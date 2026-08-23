#!/usr/bin/env python3
"""Build a page-aware, auditable search corpus for Iron Bottom Sound IV."""

from __future__ import annotations

import argparse
import hashlib
import json
import sqlite3
from pathlib import Path
from typing import Iterable


ROOT = Path(__file__).resolve().parents[3]
SOURCE_ROOT = ROOT / "resources" / "originals"
RAG_ROOT = ROOT / "resources" / "derived" / "rag"
OCR_CACHE = ROOT / "resources" / "derived" / "ocr" / "pages.jsonl"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_ocr_cache() -> dict[tuple[str, int], dict]:
    cache: dict[tuple[str, int], dict] = {}
    if not OCR_CACHE.exists():
        return cache
    for line in OCR_CACHE.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        item = json.loads(line)
        cache[(item["source"], int(item["page"]))] = item
    return cache


def pdf_pages(path: Path, ocr: dict[tuple[str, int], dict]) -> Iterable[dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    relative = path.relative_to(ROOT).as_posix()
    for index, page in enumerate(reader.pages, start=1):
        extracted = (page.extract_text() or "").strip()
        cached_item = ocr.get((relative, index), {})
        cached = str(cached_item.get("text", "")).strip()
        text = cached or extracted
        if cached:
            method = "verified_ocr" if cached_item.get("status") == "verified" else "ocr_draft"
        else:
            method = "text_layer" if extracted else "missing_ocr"
        yield {
            "source": relative,
            "page": index,
            "text": text,
            "method": method,
        }


def docx_chunks(path: Path) -> Iterable[dict]:
    from docx import Document

    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        paragraphs.append(f"表 {table_index}\n" + "\n".join(rows))
    yield {
        "source": path.relative_to(ROOT).as_posix(),
        "page": 1,
        "text": "\n".join(paragraphs),
        "method": "ooxml",
    }


def xlsx_chunks(path: Path) -> Iterable[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=False)
    for sheet in workbook.worksheets:
        lines = []
        for row in sheet.iter_rows():
            values = ["" if cell.value is None else str(cell.value) for cell in row]
            if any(values):
                lines.append(" | ".join(values))
        yield {
            "source": path.relative_to(ROOT).as_posix(),
            "page": 1,
            "text": f"工作表 {sheet.title}\n" + "\n".join(lines),
            "method": "ooxml",
        }


def image_chunk(path: Path) -> dict:
    return {
        "source": path.relative_to(ROOT).as_posix(),
        "page": 1,
        "text": path.stem.replace("-", " ").replace("_", " "),
        "method": "filename_index",
    }


def sliding_chunks(text: str, size: int = 1200, overlap: int = 200) -> Iterable[str]:
    if not text:
        yield ""
        return
    start = 0
    while start < len(text):
        yield text[start : start + size]
        if start + size >= len(text):
            break
        start += size - overlap


def build() -> dict:
    RAG_ROOT.mkdir(parents=True, exist_ok=True)
    ocr = load_ocr_cache()
    pages: list[dict] = []
    documents: list[dict] = []

    for path in sorted(p for p in SOURCE_ROOT.rglob("*") if p.is_file()):
        suffix = path.suffix.lower()
        relative = path.relative_to(ROOT).as_posix()
        document = {"source": relative, "sha256": sha256(path), "bytes": path.stat().st_size}
        if suffix == ".pdf":
            extracted = list(pdf_pages(path, ocr))
        elif suffix == ".docx":
            extracted = list(docx_chunks(path))
        elif suffix == ".xlsx":
            extracted = list(xlsx_chunks(path))
        elif suffix in {".png", ".jpg", ".jpeg"}:
            extracted = [image_chunk(path)]
        else:
            continue
        document["pages"] = len(extracted)
        document["missing_ocr_pages"] = sum(item["method"] == "missing_ocr" for item in extracted)
        document["draft_ocr_pages"] = sum(item["method"] == "ocr_draft" for item in extracted)
        document["verified_ocr_pages"] = sum(item["method"] == "verified_ocr" for item in extracted)
        documents.append(document)
        pages.extend(extracted)

    chunks = []
    tables = []
    table_terms = ("表", "Table", "命中", "修正", "结果", "穿透")
    for page in pages:
        for index, content in enumerate(sliding_chunks(page["text"]), start=1):
            chunks.append({**page, "chunk": index, "text": content})
        if any(term in page["text"] for term in table_terms):
            tables.append({k: page[k] for k in ("source", "page", "method")})

    (RAG_ROOT / "core_chunks.jsonl").write_text(
        "".join(json.dumps(item, ensure_ascii=False) + "\n" for item in chunks), encoding="utf-8"
    )
    (RAG_ROOT / "table_candidates.jsonl").write_text(
        "".join(json.dumps(item, ensure_ascii=False) + "\n" for item in tables), encoding="utf-8"
    )
    manifest = {
        "version": 1,
        "documents": documents,
        "page_count": len(pages),
        "chunk_count": len(chunks),
        "table_candidate_count": len(tables),
    }
    (RAG_ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    database = RAG_ROOT / "iron_bottom_sound_rules.sqlite3"
    if database.exists():
        database.unlink()
    connection = sqlite3.connect(database)
    connection.execute("CREATE TABLE pages (source TEXT, page INTEGER, chunk INTEGER, method TEXT, content TEXT)")
    try:
        connection.execute("CREATE VIRTUAL TABLE page_fts USING fts5(source, content, tokenize='trigram')")
    except sqlite3.OperationalError:
        connection.execute("CREATE VIRTUAL TABLE page_fts USING fts5(source, content)")
    for item in chunks:
        connection.execute(
            "INSERT INTO pages VALUES (?, ?, ?, ?, ?)",
            (item["source"], item["page"], item["chunk"], item["method"], item["text"]),
        )
        connection.execute("INSERT INTO page_fts VALUES (?, ?)", (item["source"], item["text"]))
    connection.commit()
    connection.close()
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.parse_args()
    result = build()
    print(json.dumps(result, ensure_ascii=False, indent=2))
